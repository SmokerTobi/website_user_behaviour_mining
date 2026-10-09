import os
from pptx import Presentation
from pptx.util import Inches

PPT_PATH = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'presentation', 'Website_User_Behaviour_Mining_Presentation.pptx'))

def create_presentation():
    prs = Presentation()
    
    # Slides structure
    slides_data = [
        ("Website User Behaviour Mining", "Student: Himesh Bhure\nB.Tech CSE | SIT, Nagpur\nBatch: 2024–2028 | Semester: 5 | Section: A"),
        ("Problem Statement", "- Predict e-commerce purchases based on user behavior.\n- Analyze navigation patterns and frequent categories.\n- Implement full KDD pipeline on Retailrocket data."),
        ("Dataset Description", "- Retailrocket Dataset from Kaggle.\n- Contains: Events, Item Properties, Category Tree.\n- Timestamps, Visitor IDs, Event Types (view, addtocart, transaction)."),
        ("Architecture & KDD", "- Selection: Relevant properties.\n- Cleaning: Handling missing data, invalid timestamps.\n- Transformation: Deriving 30-min sessions.\n- Data Mining: Applying ML models.\n- Presentation: Streamlit Dashboard."),
        ("Unit 1: Preprocessing", "- Timestamps parsed as Unix ms.\n- Sessions identified by 30-minute inactivity gaps.\n- Aggregated item-level events into session features (counts, duration)."),
        ("Unit 2: Concept Description", "- Generalization: Timestamps to dates.\n- Characterization: Session profiles.\n- Statistics: Comparing purchasing vs non-purchasing sessions via SQLite."),
        ("Unit 3: Classification", "- Task: Predict purchase strictly AFTER 3rd view.\n- Models: Naive Bayes, KNN, Decision Tree.\n- Evaluation: F1-score, ROC-AUC.\n- Split: By visitor to prevent leakage."),
        ("Unit 4: Clustering", "- Objective: Group sessions by behaviour.\n- Algorithms: K-Means, Agglomerative, MiniBatchKMeans.\n- Features: View count, cart count, duration, unique items.\n- Evaluated with Silhouette score."),
        ("Unit 4: Association Rules", "- Objective: Find frequently co-occurring categories in sessions.\n- Algorithms: Apriori, FP-Growth.\n- Metrics: Support, Confidence, Lift."),
        ("Unit 5: Web Mining Applications", "- Web Usage: Conversion funnel (View -> Cart -> Trans).\n- Temporal Trends: Daily event analysis."),
        ("Synthetic Demonstrations", "- Web Content: TF-IDF & Cosine Similarity on text.\n- Web Structure: PageRank on hyperlink graph.\n- Spatial Demo: Region-based visitors map.\n- (Used synthetic data for syllabus completeness)."),
        ("Dashboard Features", "- Streamlit-based interactive UI.\n- Metrics, visualizations, model outputs.\n- SQLite backend for queries."),
        ("Actual Results: Classification", "[Placeholder for F1-score chart]\n- Decision Tree typically performs robustly on categorical counts."),
        ("Actual Results: Clustering", "[Placeholder for PCA scatter plot]\n- K-Means identified distinct high-activity vs low-activity segments."),
        ("Actual Results: Association", "[Placeholder for Support/Confidence scatter]\n- Certain category pairs show high lift, indicating strong cross-sell potential."),
        ("Limitations", "- Assumption of 30-min session boundaries.\n- Excluded unmapped item categories.\n- Co-occurrence doesn't imply causation."),
        ("Future Scope", "- Real-time stream processing.\n- Deep learning for sequential pattern mining.\n- Deployment as a microservice."),
        ("Conclusion", "- Successfully implemented 5 KDD units.\n- Delivered actionable insights on user behaviour.\n- Thank You!")
    ]
    
    for title, content in slides_data:
        slide_layout = prs.slide_layouts[1] # Title and Content layout
        slide = prs.slides.add_slide(slide_layout)
        title_shape = slide.shapes.title
        body_shape = slide.placeholders[1]
        
        title_shape.text = title
        tf = body_shape.text_frame
        tf.text = content
        
    prs.save(PPT_PATH)
    print(f"Presentation saved to {PPT_PATH}")

if __name__ == '__main__':
    os.makedirs(os.path.dirname(PPT_PATH), exist_ok=True)
    create_presentation()
